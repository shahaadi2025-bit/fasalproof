#!/usr/bin/env python3
"""Adds a 'Live result' slide with REAL numbers and before/after images from a case-study run to the pitch deck.
Usage: python demo/fill_deck.py <deck.pptx> <case_output_dir> <new_deck.pptx>   (needs: pip install python-pptx)"""
import csv, json, os, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

INK, MUTED, CARD, CYAN, AMBER = RGBColor(0x0B, 0x10, 0x20), RGBColor(0x4B, 0x5B, 0x78), RGBColor(0xEA, 0xF2, 0xFF), RGBColor(0x0E, 0x74, 0x90), RGBColor(0xB4, 0x53, 0x09)
GUARD = 25.0   # losses below this are not read as damage (documented false-positive guard)

def verdict(loss, ci):
    lo = ci[0] if ci else loss
    if lo > GUARD: return "Clear drop: even the lower bound of the interval is above the 25% guard."
    if loss < GUARD: return "Below the 25% guard: not read as damage. Reported as measured."
    return "Inconclusive: the interval spans the 25% guard."

def load(case_dir):
    raw = json.load(open(os.path.join(case_dir, "raw_results.json"), encoding="utf-8")); bundle = json.load(open(os.path.join(case_dir, "signed_bundle.json"), encoding="utf-8")); p = bundle["payload"]
    rows = list(csv.DictReader(open(os.path.join(case_dir, "candidates.csv"), encoding="utf-8"))); vals = [float(r["loss_pct"]) for r in rows if r["loss_pct"]]
    return raw, p, rows, vals

def text(slide, x, y, w, h, t, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Calibri"; return tb

def card(slide, x, y, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); s.adjustments[0] = 0.08; s.fill.solid(); s.fill.fore_color.rgb = CARD; s.line.fill.background(); return s

def slide_number(slide):
    p = slide.shapes.add_textbox(Inches(9.2), Inches(5.25), Inches(0.5), Inches(0.25)).text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld")); fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"); fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr")); rpr.set("lang", "en-US"); rpr.set("sz", "1000"); etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr")).set("val", "4B5B78")
    etree.SubElement(fld, qn("a:t")).text = "‹#›"

def add_result_slide(deck_in, case_dir, deck_out, after_index=4):
    raw, p, rows, vals = load(case_dir); prs = Presentation(deck_in)
    layout = next(l for l in prs.slide_layouts if l.name == "LIGHT"); s = prs.slides.add_slide(layout)
    text(s, 0.6, 0.45, 8.8, 0.8, "What the satellites saw", 36, True, INK)
    slide_number(s)
    loss, ci = p["loss_pct"], p.get("loss_ci"); st = raw["result"].get("stats") or {}; wx = (raw.get("weather") or {}).get("clim"); ctr = p.get("controls") or []
    for i, (k, d) in enumerate((("before", p.get("scene_before")), ("after", p.get("scene_after")))):
        f = os.path.join(case_dir, k + ".png"); x = 0.6 + i * 2.75
        if os.path.exists(f): s.shapes.add_picture(f, Inches(x), Inches(1.4), Inches(2.6), Inches(2.6))
        text(s, x, 4.05, 2.6, 0.25, f"{k.capitalize()} · {d or 'n/a'}", 11, True, MUTED)
    metrics = [("Vegetation loss", f"{loss}%", f"95% CI {ci[0]}-{ci[1]}%" if ci else ""), ("Break date", str(st.get("break_date") or "n/a"), f"{st.get('offset')} days from the {p.get('loss_date')} onset" if st.get("offset") is not None else f"documented onset {p.get('loss_date')}"),
               ("Scenes used", str(p.get("scenes")), f"anomaly z {st.get('z')}" if st.get("z") is not None else ""), ("Rain vs 10 years", f"{wx['pct']:.0f}th pct" if wx else "n/a", f"{wx['cur']:.0f} mm in 3 days, median {wx['median']:.0f}" if wx else "weather unavailable")]
    for i, (a, b, c) in enumerate(metrics):
        x, y = 6.2 + (i % 2) * 1.65, 1.4 + (i // 2) * 1.3; card(s, x, y, 1.55, 1.2)
        text(s, x + 0.1, y + 0.08, 1.35, 0.25, a, 10, True, MUTED); text(s, x + 0.1, y + 0.36, 1.35, 0.4, b, 20, True, CYAN); text(s, x + 0.1, y + 0.8, 1.35, 0.35, c, 9, False, MUTED)
    text(s, 0.6, 4.4, 8.9, 0.3, verdict(loss, ci), 13, True, AMBER if loss < GUARD or (ci and ci[0] <= GUARD) else CYAN)
    rng = f"{min(vals):.1f}-{max(vals):.1f}%" if vals else "n/a"; cs = ", ".join(f"{c['season']}: {c['loss']}%" for c in ctr) or "not run"
    text(s, 0.6, 4.72, 8.9, 0.45, f"Field chosen by pre-flood canopy rule (1 of {len(rows)} candidates, loss range {rng}). Same point, earlier seasons (assumed non-event): {cs}. One field illustrates; it is not a validation.", 10, False, MUTED)
    notes = s.notes_slide.notes_text_frame; notes.text = "Real numbers from the live case-study run (demo/out). State the verdict as printed, including if it is below the guard. Mention the controls: apparent loss also appears in earlier seasons, which is why the guard exists."
    ids = prs.slides._sldIdLst; el = ids[-1]; ids.remove(el); ids.insert(after_index, el)   # place after the 'documented event' slide
    prs.save(deck_out); return deck_out

if __name__ == "__main__":
    if len(sys.argv) < 4: sys.exit(__doc__)
    print("Wrote", add_result_slide(sys.argv[1], sys.argv[2], sys.argv[3]))
