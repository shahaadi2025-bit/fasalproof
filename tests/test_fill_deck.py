import sys, pathlib, json, csv
ROOT = pathlib.Path(__file__).resolve().parents[1]; sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "demo"))
import pytest
pytest.importorskip("pptx")
import numpy as np
from pptx import Presentation
import fill_deck as fd
from ml import png_bytes
DECK = ROOT / "submission" / "FasalProof_Pitch_Deck.pptx"

def make_case(d, loss, ci, controls):
    d.mkdir(); img = png_bytes(np.full((16, 16, 3), 120, dtype="uint8")); (d / "before.png").write_bytes(img); (d / "after.png").write_bytes(img)
    json.dump({"selected": "centre", "result": {"loss_pct": loss, "stats": {"ci": ci, "z": -3.1, "break_date": "2025-08-29", "offset": 2}}, "weather": {"clim": {"cur": 140.0, "pct": 100.0, "median": 9.0, "hist": [1] * 10}}, "controls": controls}, open(d / "raw_results.json", "w"))
    json.dump({"payload": {"loss_pct": loss, "loss_ci": ci, "scenes": 9, "scene_before": "2025-08-20", "scene_after": "2025-09-06", "loss_date": "2025-08-27", "controls": [{"season": "2024-08-27", "loss": 6}]}}, open(d / "signed_bundle.json", "w"))
    with open(d / "candidates.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["point", "lat", "lon", "baseline_ndvi", "loss_pct", "scenes", "status"]); w.writerow(["centre", 1, 2, .7, loss, 9, "ok"]); w.writerow(["north", 1, 2, "", "", "", "no clear view"])

def texts(prs, i): return " ".join(sh.text_frame.text for sh in prs.slides[i].shapes if sh.has_text_frame)

def test_verdict_rules():
    assert "Clear drop" in fd.verdict(40, [30, 50]) and "not read as damage" in fd.verdict(14, [3, 20]) and "Inconclusive" in fd.verdict(35, [20, 50])

@pytest.mark.skipif(not DECK.exists(), reason="deck not built")
def test_slide_added_in_place_with_real_numbers_and_images(tmp_path):
    case = tmp_path / "case"; make_case(case, 41.5, [30.2, 52.8], []); out = tmp_path / "d.pptx"
    n0 = len(Presentation(str(DECK)).slides); fd.add_result_slide(str(DECK), str(case), str(out)); prs = Presentation(str(out)); assert len(prs.slides) == n0 + 1
    t = texts(prs, 4); assert "What the satellites saw" in t and "41.5%" in t and "95% CI 30.2-52.8%" in t and "2025-08-29" in t and "100th pct" in t and "Clear drop" in t and "1 of 2 candidates" in t and "2024-08-27: 6%" in t
    assert sum(1 for sh in prs.slides[4].shapes if sh.shape_type == 13) == 2   # two pictures
    assert "documented event" in texts(prs, 3).lower()                          # event slide sits just before the result

@pytest.mark.skipif(not DECK.exists(), reason="deck not built")
def test_weak_result_is_reported_honestly(tmp_path):
    case = tmp_path / "c2"; make_case(case, 12.0, [2.0, 21.0], []); out = tmp_path / "d2.pptx"; fd.add_result_slide(str(DECK), str(case), str(out))
    assert "not read as damage" in texts(Presentation(str(out)), 4)
